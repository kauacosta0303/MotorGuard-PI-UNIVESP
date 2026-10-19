"""Gera relatório científico conforme 'Modelo_Relatorio_Final' da UNIVESP."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
D=Document();M=json.loads((ROOT/'outputs'/'metricas.json').read_text(encoding='utf8'))
s=D.sections[0];s.page_height=Cm(29.7);s.page_width=Cm(21);s.top_margin=Cm(3);s.left_margin=Cm(3);s.bottom_margin=Cm(2);s.right_margin=Cm(2);s.header_distance=Cm(1.2)
styles=D.styles
normal=styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12);normal.font.color.rgb=RGBColor(0,0,0);normal.paragraph_format.line_spacing=1.5
normal.paragraph_format.space_after=Pt(8)
for name,size in [('Title',14),('Heading 1',12),('Heading 2',12)]:
 st=styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.font.bold=True;st.paragraph_format.space_before=Pt(5);st.paragraph_format.space_after=Pt(13)
if 'CaptionCustom' not in styles:
 st=styles.add_style('CaptionCustom',WD_STYLE_TYPE.PARAGRAPH);st.font.name='Times New Roman';st.font.size=Pt(10);st.paragraph_format.line_spacing=1;st.paragraph_format.space_after=Pt(9)

TITLE='SISTEMA INTELIGENTE PARA MONITORAMENTO E ANÁLISE DE FALHAS EM MOTORES ELÉTRICOS'
AUTHORS=[('Kauã Gabriel da Costa','23202622'),('Junio Harnold de Souza Braz','2221900'),('Julian de Almeida Santos','23203401'),('Rodrigo Marucci','2101062')]

def p(text='',indent=True):
 q=D.add_paragraph(style='Normal');q.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 if indent:q.paragraph_format.first_line_indent=Cm(1.25)
 q.add_run(text)
 return q

def center(text,sz=14,bold=False,after=12):
 q=D.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(after)
 r=q.add_run(text);r.font.name='Times New Roman';r.font.size=Pt(sz);r.bold=bold
 return q

def title1(text,new_page=False):
 if new_page:D.add_page_break()
 D.add_heading(text,1)

def title2(text,new_page=False):
 if new_page:D.add_page_break()
 D.add_heading(text,2)

def short_bullets(items):
 for t in items:
  q=D.add_paragraph(style='Normal');q.paragraph_format.left_indent=Cm(1);q.paragraph_format.first_line_indent=Cm(-.4);q.paragraph_format.space_after=Pt(3)
  q.add_run('•  '+t)

def fig(filename,cap,width=15.0):
 q=D.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.keep_with_next=True;q.paragraph_format.space_after=Pt(3)
 q.add_run().add_picture(str(ROOT/'figuras'/filename),width=Cm(width))
 q=D.add_paragraph(style='CaptionCustom');q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.add_run(cap)

def table(headers,rows,caption):
 q=D.add_paragraph(style='CaptionCustom');q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.keep_with_next=True;q.add_run(caption)
 tb=D.add_table(rows=1,cols=len(headers));tb.style='Table Grid';tb.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,h in enumerate(headers):tb.rows[0].cells[i].text=h
 for row in rows:
  cells=tb.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v)
 for idx,row in enumerate(tb.rows):
  for cell in row.cells:
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for par in cell.paragraphs:
    par.paragraph_format.space_after=Pt(2);par.paragraph_format.line_spacing=1.0
    for run in par.runs:
     run.font.name='Times New Roman';run.font.size=Pt(9.1)
     if idx==0:run.bold=True
   if idx==0:
    tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E5EAF0');tcPr.append(shd)
 D.add_paragraph().paragraph_format.space_after=Pt(4)

def codeblock(title,text):
 q=D.add_paragraph(style='CaptionCustom');q.add_run(title).bold=True;q.paragraph_format.keep_with_next=True
 for line in text.splitlines():
  a=D.add_paragraph(style='Normal');a.paragraph_format.line_spacing=1;a.paragraph_format.space_after=Pt(0);a.paragraph_format.left_indent=Cm(.8)
  run=a.add_run(line or ' ');run.font.name='Consolas';run.font.size=Pt(8.5)
 D.add_paragraph().paragraph_format.space_after=Pt(3)

def page_break():D.add_page_break()

# FRONT MATTER
center('UNIVERSIDADE VIRTUAL DO ESTADO DE SÃO PAULO',14,False)
for _ in range(3):D.add_paragraph()
for n,ra in AUTHORS:center(n+' — '+ra,14,False,5)
for _ in range(3):D.add_paragraph()
center(TITLE,14,True,20)
for _ in range(1):D.add_paragraph()
center('Vídeo de apresentação do Projeto Integrador',12,False,2)
center('[INSERIR LINK APÓS GRAVAÇÃO E PUBLICAÇÃO]',11,False,5)
for _ in range(0):D.add_paragraph()
center('Araras – SP',12,False,3);center('2026',12)

page_break();center('UNIVERSIDADE VIRTUAL DO ESTADO DE SÃO PAULO',14)
for _ in range(4):D.add_paragraph()
center(TITLE,14,True)
for _ in range(2):D.add_paragraph()
q=D.add_paragraph();q.paragraph_format.left_indent=Cm(7);q.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
q.add_run('Relatório Técnico-Científico apresentado na disciplina de Projeto Integrador para o curso de Engenharia de Computação da Universidade Virtual do Estado de São Paulo (UNIVESP).')
q=D.add_paragraph();q.paragraph_format.left_indent=Cm(7);q.add_run('Orientador: Rafael Enrique Nunes.')
for _ in range(5):D.add_paragraph()
center('Araras – SP',12,False,3);center('2026',12)

page_break()
p('COSTA, Kauã Gabriel da; BRAZ, Junio Harnold de Souza; SANTOS, Julian de Almeida; MARUCCI, Rodrigo. '+TITLE.title()+'. Relatório Técnico-Científico. Engenharia de Computação – Universidade Virtual do Estado de São Paulo. Orientador: Rafael Enrique Nunes. Polo Araras, 2026.',False)
D.add_heading('RESUMO',1)
abstract=('A manutenção de motores elétricos é relevante para a continuidade de processos industriais, pois falhas não planejadas podem comprometer a disponibilidade e elevar custos. Este estudo apresenta o MotorGuard, protótipo acadêmico para análise de registros operacionais e classificação de condições associadas a falhas mediante aprendizagem de máquina. A pesquisa utilizou 5.000 registros sintéticos com temperatura, vibração, corrente, rotação, torque e horas de operação. Foi adotada separação estratificada em treinamento (60%), validação (20%) e teste (20%), com comparação entre classificador majoritário, regressão logística e Random Forest. A seleção do algoritmo considerou a Average Precision de validação e o limiar foi escolhido pelo F2, mantendo o teste independente. A regressão logística balanceada apresentou, no teste, ROC AUC de 0,847, recall de 63,6% e precisão de 5,3%; ocorreram 126 falsos alarmes e quatro falhas não detectadas entre 1.000 registros. Foi implementada uma aplicação web dinâmica com API, gráficos, simulador e processamento CSV. O protótipo demonstra a viabilidade técnica do fluxo de dados, mas o forte desbalanceamento, o uso de dados artificiais e a ausência de validação externa impedem sua aplicação operacional. Recomenda-se coleta real, avaliação com a comunidade externa e estudo de custos de falso alarme.')
q=p(abstract,False);q.paragraph_format.line_spacing=1
q=p('PALAVRAS-CHAVE: Motores elétricos; Manutenção preditiva; Aprendizagem de máquina; Classificação de falhas; Visualização de dados.',False);q.paragraph_format.line_spacing=1

page_break();title1('LISTA DE ILUSTRAÇÕES')
for line in [
 'Figura 1 – Etapas metodológicas do Projeto Integrador',
 'Figura 2 – Amostras e distribuição das classes',
 'Figura 3 – Arquitetura funcional do sistema',
 'Figura 4 – Comparação entre algoritmos e indicadores de desempenho',
 'Figura 5 – Matriz de confusão do teste final',
 'Figura 6 – Representação conceitual do dashboard com dados do experimento'
]:p(line,False)
page_break();title1('LISTA DE TABELAS')
for line in [
 'Tabela 1 – Tecnologias e disciplinas relacionadas',
 'Tabela 2 – Dicionário de dados do experimento',
 'Tabela 3 – Descrição da divisão dos conjuntos',
 'Tabela 4 – Comparação experimental de modelos',
 'Tabela 5 – Métricas da solução selecionada',
 'Tabela 6 – Evidências de testes e estado de validação',
 'Tabela 7 – Organização dos arquivos do projeto'
]:p(line,False)
page_break();title1('SUMÁRIO')
TOC=['1 INTRODUÇÃO','2 DESENVOLVIMENTO','2.1 OBJETIVOS','2.2 JUSTIFICATIVA E DELIMITAÇÃO DO PROBLEMA','2.3 FUNDAMENTAÇÃO TEÓRICA','2.4 APLICAÇÃO DAS DISCIPLINAS ESTUDADAS NO PROJETO INTEGRADOR','2.5 METODOLOGIA','3 RESULTADOS: SOLUÇÃO FINAL','4 CONSIDERAÇÕES FINAIS','REFERÊNCIAS','APÊNDICE A — TRECHOS DE CÓDIGO','APÊNDICE B — INSTRUMENTO PARA COMUNIDADE EXTERNA','APÊNDICE C — INSTALAÇÃO E ESTRUTURA DO PROJETO']
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
PAGES=['1','2','3','4','5','7','8','11','15','16','17','18','19']
for heading,number in zip(TOC,PAGES):
 q=p('',False)
 q.paragraph_format.space_after=Pt(5)
 q.paragraph_format.tab_stops.add_tab_stop(Cm(15.3),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
 q.add_run(heading+'\t'+number)

# intro begins on own section with page numbering
sec=D.add_section(WD_SECTION_START.NEW_PAGE)
sec.top_margin=Cm(3);sec.left_margin=Cm(3);sec.bottom_margin=Cm(2);sec.right_margin=Cm(2);sec.header_distance=Cm(1.1)
sec.header.is_linked_to_previous=False
hp=sec.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=hp.add_run()
for tag,tx in [('begin',None),('instr','PAGE'),('end',None)]:
 if tag=='instr':node=OxmlElement('w:instrText');node.set(qn('xml:space'),'preserve');node.text=tx
 else:node=OxmlElement('w:fldChar');node.set(qn('w:fldCharType'),tag)
 run._r.append(node)
pg=OxmlElement('w:pgNumType');pg.set(qn('w:start'),'1');sec._sectPr.append(pg)

# Main text
D.add_heading('1 INTRODUÇÃO',1)
p('Motores elétricos estão presentes em bombas, ventiladores, compressores, transportadores e máquinas de produção. Quando um motor falha de maneira inesperada, as consequências podem incluir interrupções, manutenção emergencial e perda de produtividade. Nesse contexto, ferramentas digitais capazes de organizar informações de operação e destacar comportamentos atípicos podem auxiliar o planejamento de intervenções de manutenção.')
p('A manutenção corretiva é executada após o aparecimento de uma falha, enquanto a preventiva segue critérios periódicos de tempo ou uso. A manutenção baseada na condição busca utilizar sinais observados no equipamento para apoiar decisões. Nesse processo, variáveis como corrente, temperatura e vibração são relevantes, mas devem ser interpretadas de acordo com o tipo de motor, sua carga, condições de instalação e histórico. Por isso, a simples presença de um valor alto não é, isoladamente, comprovação de defeito.')
p('O projeto MotorGuard aborda a seguinte pergunta de pesquisa: como construir e avaliar um sistema computacional que analise variáveis operacionais e classifique registros associados a possíveis falhas, oferecendo visualizações inteligíveis para apoiar estudos de manutenção preditiva? O problema foi escolhido pela relevância técnica da disponibilidade de equipamentos e por permitir a integração de programação, ciência de dados, estatística, aprendizagem de máquina, desenvolvimento web e testes de software.')
p('O trabalho emprega uma base artificial criada para a demonstração acadêmica. Essa escolha permite exercitar a metodologia de preparação de dados, treinamento, avaliação e construção de interface sem depender de sensores físicos na fase inicial. Contudo, os registros não são séries temporais de um motor específico, e o protótipo não prevê a data de uma falha futura. O resultado fornecido é um score de classificação experimental, não uma avaliação técnica de segurança.')
p('A estrutura deste relatório segue o Modelo de Relatório Final da UNIVESP. Apresentam-se os objetivos, a justificativa, os conceitos teóricos, a aplicação dos conhecimentos do curso, a metodologia em três etapas — ouvir, criar e implementar — e os resultados mensuráveis obtidos no protótipo. Finalmente, discutem-se limitações, possíveis impactos e atividades necessárias para validação com a comunidade externa.')

title1('2 DESENVOLVIMENTO',True)
p('O desenvolvimento concentrou-se em obter uma solução demonstrável e cientificamente rastreável, com os dados, as funções de análise, o arquivo de treinamento, os testes e a interface reunidos em um único pacote. Foram adotadas práticas destinadas a reduzir vazamento de dados, documentar critérios de seleção e permitir que os resultados sejam recalculados por outro avaliador.')

title2('2.1 OBJETIVOS',True)
p('Objetivo geral: desenvolver e avaliar um protótipo web que utilize análise de dados e aprendizagem de máquina para classificar situações associadas a falhas em motores elétricos, organizando os resultados em gráficos, indicadores e recursos de simulação.')
p('Objetivos específicos:',False)
short_bullets([
 'Investigar conceitos de manutenção preditiva e variáveis operacionais relevantes para máquinas rotativas.',
 'Gerar, documentar e explorar um conjunto de dados sintéticos com seis variáveis e um rótulo de falha.',
 'Separar amostras para treinamento, validação e teste, preservando a independência do conjunto final.',
 'Comparar ao menos dois algoritmos de classificação e um classificador de referência (baseline).',
 'Selecionar limiar de classificação sem utilizar os rótulos do conjunto de teste.',
 'Implementar API REST, formulários web, gráficos interativos e importação/exportação de CSV.',
 'Criar testes automatizados e analisar tanto as falhas não detectadas quanto os falsos alarmes.',
 'Registrar limitações e organizar um protocolo de avaliação com a comunidade externa.'
])

title2('2.2 JUSTIFICATIVA E DELIMITAÇÃO DO PROBLEMA',True)
p('O motivo central da pesquisa é a necessidade de reduzir o caráter reativo das intervenções em equipamentos industriais. Em aplicações práticas, identificar indícios de degradação com antecedência pode melhorar o planejamento da manutenção. Entretanto, desenvolver um classificador é diferente de provar que ele evita falhas: essa comprovação exigiria coleta contínua, tempo de acompanhamento, comparação com históricos reais e avaliação por profissionais de manutenção.')
p('A contribuição acadêmica do MotorGuard está na integração entre engenharia e computação, permitindo reproduzir o ciclo completo de um experimento. Para fins de demonstração, o sistema organiza dados, compara algoritmos, avalia a matriz de confusão e apresenta os resultados ao usuário. O protótipo deve ser entendido como ferramenta de aprendizagem e investigação, e não como equipamento de proteção, instrumento de diagnóstico certificado ou sistema de comando de motores.')
p('A delimitação adotada compreende 5.000 registros sintéticos independentes, seis variáveis numéricas e o rótulo binário “falha”. Não estão incluídas nesta versão leituras online de CLP, sensores IoT instalados, classificação de causa-raiz, previsão de tempo restante de vida útil ou integração com sistemas industriais SCADA. Também não se dispõe de evidências documentadas de uma visita, entrevista ou validação formal com a comunidade externa, etapas que deverão ser concluídas e registradas pelo grupo antes da entrega institucional final.')
p('A investigação adota como pergunta orientadora: em um conjunto de dados rotulados, é possível obter melhor identificação de eventos classificados como falha do que a simples previsão de que todos os registros são normais, e quais compromissos entre sensibilidade e falsos alertas surgem dessa escolha?')

title2('2.3 FUNDAMENTAÇÃO TEÓRICA',True)
p('A aprendizagem supervisionada utiliza exemplos previamente rotulados para estabelecer associações entre atributos e uma classe de interesse. No caso de manutenção preditiva, o modelo pode receber variáveis operacionais e produzir um score que indica semelhança com registros rotulados como falha. Essa abordagem exige representatividade dos exemplos; relações incorporadas artificialmente ao gerador não necessariamente se repetem em máquinas reais.')
p('A regressão logística estima uma função sigmoide da combinação linear dos atributos e pode oferecer uma referência interpretável para classificação binária. O algoritmo Random Forest combina múltiplas árvores de decisão treinadas com aleatorização, conforme apresentado por Breiman (2001). Ambos são métodos relevantes em experimentos tabulares, mas devem ser comparados por métricas apropriadas e com dados separados para avaliação.')
p('Em bases fortemente desbalanceadas, a acurácia é insuficiente para avaliação. Se apenas 1% dos exemplos é positivo, uma regra que sempre responde “sem falha” pode apresentar 99% de acerto e, ainda assim, deixar passar todos os casos de interesse. A precisão mede a proporção de alertas que corresponde a rótulos positivos, enquanto o recall mede a proporção de casos positivos que foram identificados. A Average Precision (AP) sintetiza a relação entre precisão e recall em diferentes limiares e pode ser mais informativa que a acurácia quando falhas são raras.')
p('Para reduzir o risco de conclusões otimistas, a documentação do scikit-learn recomenda separar os conjuntos antes de aplicar transformações que aprendam parâmetros, ajustando pré-processamento e modelo exclusivamente aos dados de treinamento (SCIKIT-LEARN DEVELOPERS, 2026). No MotorGuard, essa recomendação foi incorporada a pipelines, e a escolha do algoritmo e do limiar ocorreu na validação. O teste final foi reservado para uma avaliação única, com registro explícito de verdadeiros e falsos positivos e negativos.')
p('Também foi consultada a documentação do AI4I 2020 Predictive Maintenance Dataset, disponibilizado pela UCI (2020), como possível caminho de ampliação do estudo. Trata-se de um benchmark sintético sobre manutenção de máquinas e NÃO foi a fonte dos números deste relatório. A base efetivamente utilizada é a simulação local disponibilizada no pacote, acompanhada de seu gerador com seed 42. Essa distinção é importante para assegurar rastreabilidade e evitar atribuição indevida de resultados à UCI.')

title2('2.4 APLICAÇÃO DAS DISCIPLINAS ESTUDADAS NO PROJETO INTEGRADOR',True)
p('O projeto articula conhecimentos de diferentes componentes curriculares de Engenharia de Computação. A relação não se limita à teoria: os conteúdos aparecem na preparação dos dados, no treinamento dos algoritmos, na organização do código, na aplicação web e nos testes. A tabela seguinte apresenta essa correspondência de forma verificável.')
table(['Área / disciplina','Conteúdo aplicado','Evidência técnica'],[
 ('Algoritmos e programação','Python; funções; tratamento de listas e dados','ml/train.py; data/gerar_base.py'),
 ('Probabilidade e estatística','Amostragem; prevalência; métricas; separação de conjuntos','outputs/metricas.json; fig. 2, 4 e 5'),
 ('Inteligência artificial','Regressão logística; Random Forest; classificação supervisionada','Pipelines do scikit-learn'),
 ('Desenvolvimento web','HTML, CSS, JavaScript; interface dinâmica; formulários','app/static/index.html, style.css e app.js'),
 ('Redes e APIs','Requisições HTTP; JSON; serviço REST','app/main.py; rota /docs'),
 ('Engenharia de software','Validação de entrada; organização modular; testes','tests/test_api.py; tests/test_training.py'),
 ('Instrumentação e IoT (perspectiva)','Temperatura, corrente e vibração como sinais de condição','Dicionário de variáveis; integração futura')
],'Tabela 1 – Relação entre conteúdos acadêmicos e evidências do projeto. Fonte: elaborado pelo grupo.')
p('É importante distinguir os recursos implementados das propostas futuras. A API e os testes estão presentes no pacote e podem ser executados. A comunicação com sensores IoT, a aquisição por gateway e o armazenamento contínuo em banco de dados ainda são possibilidades de evolução, e não funcionalidades concluídas. Os conteúdos dessas áreas contribuem para orientar requisitos futuros, mas não devem ser contabilizados como implementação efetiva.')

title2('2.5 METODOLOGIA',True)
p('A pesquisa possui natureza aplicada e abordagem quantitativa para o experimento computacional, com investigação exploratória do contexto. O trabalho foi organizado segundo as três etapas indicadas no modelo da UNIVESP: ouvir e interpretar o contexto, criar/prototipar e implementar/testar. Além da solução digital, são descritas as evidências disponíveis e as lacunas que ainda dependem de interação real com a comunidade externa.')
fig('fig02_metodologia.png','Figura 1 – Etapas metodológicas do Projeto Integrador. Fonte: elaboração própria.')
p('Etapa 1 — Ouvir e interpretar o contexto. O grupo definiu como problema motivador a ocorrência de paradas inesperadas de motores. Foram levantadas variáveis de interesse e consultada literatura de classificação e avaliação. Não se dispõe, até esta versão, de registros documentais de entrevistas, visitas ou validação externa. Para concluir essa etapa institucional, o Apêndice B oferece roteiro de entrevista e formulário de avaliação a serem preenchidos após interação real com profissionais.')
p('Etapa 2 — Criar e prototipar. A base sintética foi construída por um gerador Python com seed 42, com 5.000 observações. As distribuições de temperatura, rotação, torque, vibração, corrente e horas foram definidas matematicamente, e a classe de falha foi sorteada a partir de uma função logística de variáveis selecionadas. Assim, o rótulo não resulta de uma pane física: ele decorre de regras probabilísticas estabelecidas pelo gerador, presentes em data/gerar_base.py.')
table(['Variável','Unidade','Papel','Origem no protótipo'],[
 ('temperatura_c','°C','Entrada','Simulação numérica'),
 ('vibracao_mm_s','mm/s','Entrada','Simulação numérica'),
 ('corrente_a','A','Entrada','Simulação numérica'),
 ('rotacao_rpm','rpm','Entrada','Simulação numérica'),
 ('torque_nm','N·m','Entrada','Simulação numérica'),
 ('horas_operacao','h','Entrada','Simulação numérica'),
 ('falha','0 ou 1','Rótulo, não entrada','Sorteio condicionado a função logística')
],'Tabela 2 – Dicionário de dados utilizado no experimento. Fonte: elaboração própria.')
p('O arquivo CSV foi separado por amostragem estratificada e seed 42, preservando aproximadamente a proporção de positivos em cada subconjunto. Os 3.000 exemplos de treinamento alimentaram os modelos; os 1.000 de validação serviram para a escolha do algoritmo e do limiar; e os 1.000 finais permaneceram fora de todas as decisões até a avaliação final. Como os registros não representam uma sequência temporal real, essa separação não constitui validação cronológica.')
fig('fig03_distribuicao.png','Figura 2 – Divisão experimental e desbalanceamento da base sintética. Fonte: resultados reproduzidos pelo código.')
table(['Conjunto','Registros','Função'],[('Treinamento','3.000','Ajuste dos parâmetros dos modelos'),('Validação','1.000','Seleção por AP e limiar por F2'),('Teste','1.000','Avaliação final, sem ajuste do modelo')],'Tabela 3 – Divisão dos subconjuntos experimentais. Fonte: elaboração própria.')
p('Foram treinadas três alternativas: baseline de classe majoritária, regressão logística com ponderação de classes e Random Forest com balanceamento por classe. A imputação de valores e a padronização, quando necessária, foram encapsuladas em Pipeline para evitar aplicar transformações aprendidas sobre a validação ou o teste. Escolheu-se o algoritmo com maior Average Precision de validação. O limiar de decisão foi definido por busca em grade de 0,03 a 0,95, maximizando F2 na validação, que confere maior peso ao recall.')
p('Etapa 3 — Implementar e testar. O modelo foi serializado com joblib e conectado à API FastAPI, responsável por validar entradas, calcular scores, retornar JSON e processar CSV. A interface utiliza HTML, CSS, JavaScript e Plotly hospedado localmente no próprio pacote, com páginas de visão geral, análise exploratória, simulador, métricas, importação e metodologia. Os testes automatizados verificam o esquema do dataset, o cálculo de métricas, o endpoint de saúde, a predição e a rejeição de entradas inválidas.')
fig('fig01_arquitetura.png','Figura 3 – Arquitetura do sistema implementado. Fonte: elaboração própria.')
p('O fluxo foi testado por chamadas diretas à API e por testes de integração automatizados. A apresentação visual incorporada ao relatório é um mockup conceitual alinhado ao layout e aos dados do experimento; não é uma fotografia de um sistema industrial em operação. Eventuais avaliações de usabilidade de profissionais externos deverão ser realizadas e anexadas posteriormente.')

title1('3 RESULTADOS: SOLUÇÃO FINAL',True)
p('O resultado técnico entregue consiste em um protótipo reproduzível que integra uma base de 5.000 registros simulados, um gerador de dados, um script de treinamento e comparação de modelos, um classificador salvo, uma API REST, uma interface web dinâmica e testes automatizados. O sistema é executado localmente e não utiliza aquisição online de sensores. Seus resultados caracterizam classificação de registros, e não previsão temporal comprovada de falhas futuras.')
p('Na seleção experimental, a regressão logística balanceada obteve AP de validação de 0,128, superior ao Random Forest (0,030) e ao baseline (0,011). Por essa razão, a versão final deste protótipo utiliza a regressão logística, mesmo que a primeira proposta de projeto destacasse Random Forest. Essa decisão é coerente com o critério previamente estabelecido e demonstra a importância de comparar alternativas.')
table(['Algoritmo','AP validação','ROC AUC validação','Limiar escolhido'],[
 (x['modelo'],f"{x['ap_validacao']:.3f}",f"{x['roc_auc_validacao']:.3f}",f"{x['limiar_validacao']:.2f}") for x in M['modelos_comparados']
],'Tabela 4 – Desempenho dos modelos avaliados na validação. Fonte: outputs/metricas.json.')
fig('fig05_metricas.png','Figura 4 – Resultados da comparação e do teste final. Fonte: dados experimentais reproduzíveis.')
p('O conjunto de teste contém 1.000 registros, entre os quais 11 rotulados com falha. O limiar de 0,66 foi determinado antes da avaliação final. O modelo sinalizou corretamente 7 falhas e deixou 4 sem sinalização; simultaneamente, gerou 126 alarmes em registros rotulados como normais. A precisão calculada foi 5,26%, o recall foi 63,64% e a Average Precision no teste foi 0,045. A ROC AUC de 0,847 revela capacidade de ordenação, mas não anula o excesso de falsos positivos.')
table(['Indicador','Valor','Interpretação'],[
 ('ROC AUC','0,847','Capacidade geral de distinguir rótulos ao variar limiar'),
 ('Average Precision','0,045','Qualidade do ranqueamento sob classe rara'),
 ('Acurácia','87,0%','Não deve ser utilizada isoladamente'),
 ('Precisão','5,26%','Poucos alertas positivos correspondem a falhas rotuladas'),
 ('Recall','63,64%','7 de 11 registros positivos identificados'),
 ('F1','0,097','Média harmônica de precisão e recall'),
 ('VP / FN / FP / VN','7 / 4 / 126 / 863','Contagens da matriz de confusão')
],'Tabela 5 – Métricas no conjunto de teste intocado. Fonte: outputs/metricas.json.')
fig('fig04_matriz_confusao.png','Figura 5 – Matriz de confusão do teste final (rótulos sintéticos). Fonte: elaboração própria.')
p('A comparação com o baseline reforça a relevância do desbalanceamento. Classificar todos os registros como normais produz 98,9% de acurácia, mas não identifica nenhuma das 11 falhas. O modelo selecionado reduz essa cegueira, embora com elevado custo de falsos alertas. A queda de AP entre validação (0,128) e teste (0,045), associada ao pequeno número de positivos, indica instabilidade e exige novos experimentos e dados de maior representatividade.')
p('O dashboard foi estruturado para apoiar a compreensão do experimento. A visão geral apresenta dados descritivos; a exploração permite selecionar e comparar variáveis; o simulador recebe seis parâmetros e informa um score experimental; a aba científica exibe métricas e matriz de confusão; e a importação CSV realiza predições em lote. A interface não exibe status de motores físicos, porque cada linha representa um registro artificial independente.')
fig('fig06_mockup_dashboard.png','Figura 6 – Representação conceitual do dashboard, com indicadores derivados do experimento. Fonte: elaborado pelo grupo; não é captura de sensores reais.')
p('Além dos recursos funcionais, foram produzidos testes automatizados. O comando “python -m pytest tests -q” apresentou dez testes aprovados no ambiente de desenvolvimento. O serviço web respondeu a requisições em /api/health e gerou predições e relatórios de métricas, conforme os contratos da API. A validação por navegador e entrevistas de campo não foi considerada concluída nesta versão.')
table(['Verificação','Evidência','Situação'],[
 ('Treinamento reproduzível','python -m ml.train','Executado'),
 ('Regeneração idêntica do CSV','data/gerar_base.py; seed 42','Verificada'),
 ('Testes unitários/integrados','pytest: 10 testes','Aprovado'),
 ('Resposta API de saúde','GET /api/health','Verificada'),
 ('Validação dos campos','POST /api/predict e /api/batch','Verificada por testes'),
 ('Validação em motores reais','Coleta de sensor / histórico físico','Não realizada'),
 ('Feedback da comunidade externa','Entrevista e formulário assinado','Pendente')
],'Tabela 6 – Evidências efetivas e etapas ainda não validadas. Fonte: registros do projeto.')
p('Quanto às melhorias entre a solução inicial e esta versão, houve a substituição de um painel introdutório por uma aplicação com API e várias telas, divisão em três conjuntos, baseline comparativo, busca de limiar, upload de CSV, validações e testes. Essas mudanças decorreram do refinamento técnico do próprio desenvolvimento; não devem ser atribuídas a comentários de uma comunidade externa antes que tais comentários sejam efetivamente coletados e documentados.')

title1('4 CONSIDERAÇÕES FINAIS',True)
p('O MotorGuard atingiu o objetivo técnico de demonstrar um processo completo de análise de dados e classificação supervisionada, integrado a uma aplicação web com gráficos, simulações, importação de dados e documentação. O pacote permite a reprodução dos resultados por meio de códigos organizados e oferece um recurso didático para discutir manutenção preditiva no contexto de Engenharia de Computação.')
p('Os resultados também evidenciam limites relevantes. Apesar de um recall de 63,64%, a precisão de 5,26% representa grande volume de falsos alarmes; além disso, quatro dos onze registros positivos do conjunto de teste não foram identificados. A baixa prevalência de falhas limita a estabilidade das métricas. O modelo foi treinado e testado apenas com registros gerados por funções estatísticas escolhidas para a simulação. Portanto, não há evidência de que os scores forneçam risco real ou permitam evitar panes em motores industriais.')
p('A contribuição imediata é acadêmica: organização do método científico, programação, comparação de algoritmos, interpretação de métricas e apresentação clara de resultados. Em relação à comunidade externa, o impacto permanece potencial, pois ainda não existe documentação de adoção, validação ou retorno formal. O grupo deverá realizar contato efetivo com profissionais ou organização participante, demonstrar o protótipo, registrar comentários e atualizar este relatório com evidências verificáveis.')
p('Como continuidade, recomenda-se empregar dados físicos obtidos sob autorização, incluir séries temporais e contexto de operação, avaliar calibração probabilística, discutir custos de falsos positivos e negativos, testar diferentes equipamentos, adotar validação temporal e verificar possíveis efeitos de mudança do processo. Essas ações seriam necessárias para evoluir de um demonstrador acadêmico para uma solução de apoio à manutenção avaliada em campo.')

title1('REFERÊNCIAS',True)
refs=[
 'BREIMAN, Leo. Random Forests. Machine Learning, v. 45, p. 5–32, 2001. DOI: https://doi.org/10.1023/A:1010933404324.',
 'SCIKIT-LEARN DEVELOPERS. Common pitfalls and recommended practices. Disponível em: https://scikit-learn.org/stable/common_pitfalls.html. Acesso em: 18 out. 2026.',
 'SCIKIT-LEARN DEVELOPERS. Precision-Recall. Disponível em: https://scikit-learn.org/stable/auto_examples/model_selection/plot_precision_recall.html. Acesso em: 18 out. 2026.',
 'UCI MACHINE LEARNING REPOSITORY. AI4I 2020 Predictive Maintenance Dataset. 2020. DOI: https://doi.org/10.24432/C5HS5C. Disponível em: https://archive.ics.uci.edu/dataset/601/ai4i+2020+predictive+maintenance+dataset. Acesso em: 18 out. 2026. Nota: referência de benchmark para possível estudo futuro; não é a base experimental deste relatório.',
 'FASTAPI. FastAPI Documentation. Disponível em: https://fastapi.tiangolo.com/. Acesso em: 18 out. 2026.',
 'PLOTLY. Plotly JavaScript Open Source Graphing Library. Disponível em: https://plotly.com/javascript/. Acesso em: 18 out. 2026.',
 'UNIVERSIDADE VIRTUAL DO ESTADO DE SÃO PAULO (UNIVESP). Modelo de Relatório Final do Projeto Integrador. Documento disponibilizado no AVA, 2026.'
]
for r in refs:
 q=p(r,False);q.paragraph_format.line_spacing=1;q.paragraph_format.space_after=Pt(13)

title1('APÊNDICE A — TRECHOS DE CÓDIGO',True)
p('Os códigos abaixo são excertos representativos das rotinas executadas no protótipo. A íntegra dos arquivos Python, HTML, CSS e JavaScript está disponível no pacote digital MotorGuard_UNIVESP_Final.zip. Todas as rotinas apresentadas podem ser reproduzidas sem acesso a equipamentos industriais.')
codeblock('Código A1 – Separação estratificada, sem utilizar o teste na seleção:', '''trtmp, teste = train_test_split(indices, test_size=.20,
    stratify=y, random_state=42)
treino, validacao = train_test_split(trtmp, test_size=.25,
    stratify=y.iloc[trtmp], random_state=42)
# 60% treino / 20% validacao / 20% teste''')
codeblock('Código A2 – Pipeline de regressão logística balanceada:', '''pipeline = make_pipeline(
    SimpleImputer(strategy='median'),
    StandardScaler(),
    LogisticRegression(class_weight='balanced',
                       random_state=42, max_iter=2000)
)
pipeline.fit(X.iloc[treino], y.iloc[treino])''')
codeblock('Código A3 – Seleção do limiar exclusivamente na validação:', '''limiares = np.round(np.linspace(.03, .95, 93), 3)
valores = [fbeta_score(y_val, probs_val >= t,
                       beta=2, zero_division=0)
           for t in limiares]
limiar = float(limiares[np.argmax(valores)])''')
codeblock('Código A4 – Predição por endpoint HTTP:', '''@app.post('/api/predict')
def predict(item: MotorInput):
    dados = pd.DataFrame([item.model_dump()])
    score = float(predict_df(dados)[0])
    return {'score': round(score, 6),
            'classificacao': 'ATENCAO' if score >= threshold
                               else 'SEM SINALIZACAO'}''')
p('O exemplo A3 é uma simplificação do código utilizado; a função completa select_threshold também define desempate entre limiares. Da mesma forma, o endpoint real acrescenta mensagens de segurança e valida os intervalos de entrada. Os arquivos integrais do pacote são a referência executável do projeto.')

title1('APÊNDICE B — INSTRUMENTO PARA COMUNIDADE EXTERNA',True)
p('Este instrumento foi preparado para coletar evidências reais, atendendo à exigência do modelo final de relatar o contexto, a implementação e o retorno da comunidade externa. Não houve preenchimento fictício. Deve ser aplicado pelo grupo e anexado após a interação com a organização participante.')
short_bullets([
 'Organização ou comunidade participante: ______________________________________',
 'Responsável entrevistado e função: _________________________________________',
 'Data, local e forma de contato: _____________________________________________',
 'Principais problemas de manutenção observados: _____________________________',
 'Quais dados de motores são disponíveis (vibração, corrente, temperatura etc.)?',
 'A interface é compreensível? O que dificulta a interpretação dos resultados?',
 'Que limites mínimos de precisão/recall seriam necessários para o contexto?',
 'Que melhorias o participante sugere? Há autorização para descrever a visita?',
 'Quais ajustes foram realizados depois do feedback? __________________________'
])
p('A documentação complementar poderá incluir termo de autorização, registros da visita, capturas da solução demonstrada e síntese da devolutiva, sempre respeitando a confidencialidade de dados industriais e de participantes.')

title1('APÊNDICE C — INSTALAÇÃO E ESTRUTURA DO PROJETO',True)
p('O pacote pode ser executado em computador Windows com Python instalado, sem necessidade de hardware IoT. Uma vez instaladas as dependências, a interface JavaScript utiliza uma biblioteca Plotly incluída localmente. A instalação inicial das dependências poderá requerer conexão com a internet.')
codeblock('Código C1 – Comandos de execução (Windows PowerShell):', '''py -3 -m venv .venv
.venv\\Scripts\\python.exe -m pip install -r requirements.txt
.venv\\Scripts\\python.exe -m ml.train
.venv\\Scripts\\python.exe -m pytest tests -q
.venv\\Scripts\\python.exe -m uvicorn app.main:app \\
    --host 127.0.0.1 --port 8000''')
p('Após a inicialização, abrir http://127.0.0.1:8000 no navegador. A documentação interativa das rotas está disponível em http://127.0.0.1:8000/docs. O arquivo INICIAR_WINDOWS.bat automatiza a criação de ambiente, instalação e inicialização em Windows.')
table(['Arquivo ou diretório','Descrição'],[
 ('data/gerar_base.py','Gerador reproduzível de registros sintéticos'),
 ('data/motores_sinteticos.csv','Base da pesquisa: 5.000 registros'),
 ('ml/train.py','Treino, seleção, avaliação e exportação'),
 ('models/motorguard.joblib','Modelo e limiar experimentais salvos'),
 ('outputs/metricas.json','Resultados e configurações verificáveis'),
 ('app/main.py','API REST e validação de dados'),
 ('app/static/*','Interface HTML/CSS/JavaScript e Plotly local'),
 ('tests/*','Testes automatizados de dados e API'),
 ('figuras/*','Ilustrações, análises e mockup do dashboard'),
 ('documentos/*','Relatório final em Word e apresentação')
],'Tabela 7 – Organização do pacote de entrega. Fonte: elaboração própria.')
p('Pendências para entrega institucional: inserir o link do vídeo gravado, validar a interação com a comunidade externa, registrar o feedback, revisar datas e demais dados institucionais no AVA, confirmar as referências e atualizar a seção de impacto com fatos comprovados.')

D.core_properties.title=TITLE.title();D.core_properties.author='Grupo Projeto Integrador — UNIVESP';D.core_properties.subject='Relatório final acadêmico, versão técnica para revisão — MotorGuard'
out=ROOT/'documentos'/'Relatorio_Final_MotorGuard_UNIVESP.docx'
D.save(out)
print(out,'paragraphs',len(D.paragraphs),'tables',len(D.tables),'images',len(D.inline_shapes))
