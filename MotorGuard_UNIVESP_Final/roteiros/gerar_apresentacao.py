"""Apresentação acadêmica consistente com os dados verificáveis do relatório."""
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
R=Presentation();R.slide_width=Inches(13.333);R.slide_height=Inches(7.5)
NAVY=RGBColor(10,23,39); PANEL=RGBColor(17,37,56); GREEN=RGBColor(126,227,183); WHITE=RGBColor(243,250,255); MUTED=RGBColor(161,183,203); ORANGE=RGBColor(248,171,117)
def rect(sl,x,y,w,h,fill,radius=False):
 sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=fill;sh.line.fill.background();return sh

def txt(sl,text,x,y,w,h,size=20,color=WHITE,bold=False,align=None):
 sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=Pt(0);tf.margin_top=Pt(0)
 para=tf.paragraphs[0];para.text=text;para.font.name='Aptos';para.font.size=Pt(size);para.font.color.rgb=color;para.font.bold=bold
 if align:para.alignment=align
 return sh

def base(kicker,title,subtitle=''):
 sl=R.slides.add_slide(R.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=NAVY
 rect(sl,.0,0,.12,7.5,GREEN)
 txt(sl,kicker.upper(),.6,.42,11,.3,size=12,color=GREEN,bold=True)
 txt(sl,title,.6,.95,12,1.0,size=32,bold=True)
 if subtitle:txt(sl,subtitle,.64,1.92,12,.9,size=16,color=MUTED)
 txt(sl,'UNIVESP · MOTORGUARD · 2026',.65,7.14,10,.24,size=10,color=MUTED)
 txt(sl,str(len(R.slides)).zfill(2),12.15,7.1,.6,.25,size=10,color=MUTED,align=PP_ALIGN.RIGHT)
 return sl

def points(sl,items,x=.7,y=2.7,width=11.8,size=23,spacing=.83):
 for i,item in enumerate(items):
  rect(sl,x,y+i*spacing+.11,.09,.09,GREEN)
  txt(sl,item,x+.25,y+i*spacing,width-.35,.7,size=size)

def picture(sl,f,x=1.3,y=2.15,w=10.9):
 sl.shapes.add_picture(str(ROOT/'figuras'/f),Inches(x),Inches(y),width=Inches(w))

s=base('projeto integrador · engenharia de computação','MOTORGUARD','Sistema inteligente para análise experimental de falhas em motores elétricos')
points(s,['Kauã Gabriel da Costa · Junio Harnold de Souza Braz','Julian de Almeida Santos · Rodrigo Marucci','Orientador: Rafael Enrique Nunes','Relatório final técnico — versão para revisão'],y=2.8,size=20,spacing=.75)
s=base('desafio industrial','Falhas inesperadas e decisões de manutenção','A pergunta: é possível identificar condições relacionadas a falhas a partir de dados?')
points(s,['Paradas não programadas elevam custos e indisponibilidade','Dados operacionais podem apoiar análises e priorização','Identificar uma classe NÃO equivale a prever a data de falha'],size=22)
s=base('escopo e honestidade científica','Dados sintéticos, não sensores reais','5.000 registros experimentais, 54 rotulados com falha (1,08%).')
picture(s,'fig03_distribuicao.png',x=1.0,y=2.4,w=11.0)
s=base('dados e variáveis','Seis entradas → um rótulo de falha')
points(s,['Temperatura (°C) · vibração (mm/s) · corrente (A)','Rotação (rpm) · torque (N·m) · horas de operação','“Falha” é apenas o rótulo binário artificial usado para treinar'],size=21)
s=base('arquitetura implementada','Do dataset até o navegador')
picture(s,'fig01_arquitetura.png',x=.55,y=2.5,w=12.2)
s=base('rigor metodológico','Treino, validação e teste separados')
points(s,['60% treinamento · 20% validação · 20% teste','Pré-processamento dentro do pipeline, sem vazamento','Modelo selecionado pela Average Precision na validação','Limiar selecionado por F2 na validação'],size=20,spacing=.77)
s=base('comparação de modelos','A regressão logística venceu na validação')
picture(s,'fig05_metricas.png',x=.55,y=2.3,w=12.3)
s=base('resultado no teste final','7 positivos detectados · 4 não detectados','O conjunto final tem somente 11 positivos e 1.000 registros.')
picture(s,'fig04_matriz_confusao.png',x=3.25,y=2.15,w=6.6)
s=base('interpretação crítica','Não basta apresentar uma ROC AUC alta')
points(s,['Recall: 63,6% · precisão: 5,26%','126 falsos alarmes em 1.000 registros testados','Baseline: 98,9% de acurácia, mas 0% de recall','O desempenho não é validado para motores reais'],size=21,spacing=.78)
s=base('interatividade','Dashboard + API + simulador + CSV')
picture(s,'fig06_mockup_dashboard.png',x=.7,y=2.16,w=11.95)
s=base('testes e reprodutibilidade','Código, dados, modelo e testes no pacote')
points(s,['python data/gerar_base.py → reproduz a base sintética','python -m ml.train → treina e registra métricas','python -m pytest tests -q → 10 testes aprovados','python -m uvicorn app.main:app → inicia o dashboard'],size=21,spacing=.78)
s=base('etapas de pesquisa','Ouvir · criar · implementar')
picture(s,'fig02_metodologia.png',x=1.0,y=2.18,w=11.5)
s=base('fechamento','Contribuições e próximos passos')
points(s,['Protótipo acadêmico completo, demonstrável e auditável','Coletar dados reais, avaliar deriva e melhorar precisão','Documentar visita e feedback de uma comunidade externa','Gravar vídeo e atualizar relatório antes da entrega'],size=21,spacing=.8)
output=ROOT/'documentos'/'Apresentacao_MotorGuard_UNIVESP.pptx';R.save(output);print(output,len(R.slides),'slides')
